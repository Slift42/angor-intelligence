"""Génère le classeur Excel de revue de l'analyse des risques (tools/build_risk_workbook.py).

Le classeur reprend les VALEURS ACTUELLES lues dans le code et la configuration (config/*.json, veille/*.py,
docs/gonogo.js) et propose, à côté, des colonnes jaunes où l'analyste note ses propositions, ainsi que des
simulateurs (formules Excel) qui reproduisent les calculs de l'outil.
Usage : python tools/build_risk_workbook.py [chemin_sortie.xlsx] [--force] [--lang en] [--missing]
(par défaut revue/Analyse_des_risques_Angor.xlsx, ou revue/Angor_risk_analysis_review.xlsx en anglais ; un fichier
existant est conservé et le nouveau prend la date, sauf --force ; traduction : tools/workbook_i18n.py)
Le classeur rempli est ensuite relu par le développeur, qui reporte les décisions dans config/ et le code
(voir l'onglet « Lisez-moi »).
"""
import csv
import json
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
from veille import model, press, quality  # noqa: E402
from veille.dedupe import FAMILIES  # noqa: E402
from veille.risk import DEFAULTS as RISK_DEFAULTS  # noqa: E402

# Sortie par défaut dans revue/ (ignoré par Git : le dépôt est public, les choix de méthode de l'analyste ne le sont pas).
# Un classeur existant n'est jamais écrasé (il peut contenir les annotations de l'analyste) : on ajoute la date.
# --lang en : version anglaise (traduite après construction, voir tools/workbook_i18n.py) ;
# --missing : écrit aussi la liste des textes sans traduction dans revue/traductions_manquantes.json.
LANG = sys.argv[sys.argv.index("--lang") + 1] if "--lang" in sys.argv else "fr"
_args = [a for i, a in enumerate(sys.argv[1:], 1)
         if not a.startswith("--") and sys.argv[i - 1] != "--lang"]
_default = "Analyse_des_risques_Angor.xlsx" if LANG == "fr" else "Angor_risk_analysis_review.xlsx"
OUT = Path(_args[0]) if _args else ROOT / "revue" / _default
if OUT.exists() and "--force" not in sys.argv:
    OUT = OUT.with_name(f"{OUT.stem}_{date.today():%Y-%m-%d}{OUT.suffix}")
OUT.parent.mkdir(parents=True, exist_ok=True)
CFG = ROOT / "config"
RISK = {**RISK_DEFAULTS, **json.loads((CFG / "risk.json").read_text(encoding="utf-8"))}
SETTINGS = json.loads((CFG / "settings.json").read_text(encoding="utf-8"))
HEALTH = json.loads((CFG / "health.json").read_text(encoding="utf-8"))
NOTES = json.loads((CFG / "city_notes.json").read_text(encoding="utf-8"))
SOURCES = json.loads((CFG / "sources.json").read_text(encoding="utf-8"))["sources"]
OUTLETS = json.loads((CFG / "press_outlets.json").read_text(encoding="utf-8"))["countries"]
_c = (ROOT / "docs" / "data" / "countries.js").read_text(encoding="utf-8")
COUNTRIES = sorted((f["properties"] for f in json.loads(_c[_c.index("=") + 1:].strip().rstrip(";"))["features"]),
                   key=lambda p: p["name_fr"])
CNAME = {p["iso2"]: p["name_fr"] for p in COUNTRIES}
GOLD = list(csv.DictReader(open(ROOT / "tests" / "gold_tri.csv", encoding="utf-8")))

# ------------------------------------------------------------------ styles
NAVY, GOLD_C = "0E1B2C", "B08D57"
F = "Arial"
FONT = Font(name=F, size=10)
BOLD = Font(name=F, size=10, bold=True)
TITLE = Font(name=F, size=16, bold=True, color=NAVY)
SUB = Font(name=F, size=10, italic=True, color="555555")
H2 = Font(name=F, size=12, bold=True, color=NAVY)
HEAD = Font(name=F, size=10, bold=True, color="FFFFFF")
INPUT_FONT = Font(name=F, size=10, color="0000FF")
HEAD_FILL = PatternFill("solid", fgColor=NAVY)
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")   # jaune clair : à modifier
CALC_FILL = PatternFill("solid", fgColor="EEF1F4")    # gris : calcul automatique
CUR_FILL = PatternFill("solid", fgColor="FFFFFF")
NOTE_FILL = PatternFill("solid", fgColor="F4EEE3")
THIN = Side(style="thin", color="D0D6DD")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)

wb = Workbook()
wb.remove(wb.active)
TABS = []   # (nom, ce qu'on y décide)


def sheet(name, title, intro, decide):
    ws = wb.create_sheet(name)
    TABS.append((name, decide))
    ws.sheet_view.showGridLines = False
    ws["A1"] = title
    ws["A1"].font = TITLE
    ws["A2"] = intro
    ws["A2"].font = SUB
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A2:H2")
    ws.row_dimensions[2].height = 48
    ws["A3"] = "Légende : cellules jaunes à texte bleu = à modifier par vous · cellules grises = calcul automatique · le reste = valeur actuelle de l'outil."
    ws["A3"].font = Font(name=F, size=9, color="888888")
    return ws


def widths(ws, *w):
    for i, x in enumerate(w, 1):
        ws.column_dimensions[get_column_letter(i)].width = x


def header(ws, row, labels, col=1):
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=lab)
        c.font, c.fill, c.alignment, c.border = HEAD, HEAD_FILL, CENTER, BOX
    ws.row_dimensions[row].height = 30


def put(ws, row, col, value, kind="cur", fmt=None, bold=False):
    """kind : cur (valeur actuelle), in (à modifier), calc (formule), note."""
    c = ws.cell(row=row, column=col, value=value)
    c.border = BOX
    c.alignment = WRAP
    if kind == "in":
        c.font, c.fill = INPUT_FONT, INPUT_FILL
    elif kind == "calc":
        c.font, c.fill = (BOLD if bold else FONT), CALC_FILL
    elif kind == "note":
        c.font, c.fill = Font(name=F, size=9, italic=True, color="555555"), NOTE_FILL
    else:
        c.font = BOLD if bold else FONT
    if fmt:
        c.number_format = fmt
    return c


def section(ws, row, text):
    ws.cell(row=row, column=1, value=text).font = H2
    return row + 1


def dv_list(ws, rng, values, prompt=None):
    dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
    if prompt:
        dv.promptTitle, dv.prompt, dv.showInputMessage = "Choix", prompt, True
    ws.add_data_validation(dv)
    dv.add(rng)


def finish(ws, freeze=None, filt=None):
    if freeze:
        ws.freeze_panes = freeze
    if filt:
        ws.auto_filter.ref = filt
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0


TODAY = date.today().strftime("%d/%m/%Y")
LEVEL_NAMES = {k: v["fr"] for k, v in model.RISK_LEVELS.items()}
SEV_NAMES = {k: v["fr"] for k, v in model.SEVERITY.items()}

# ================================================================== Lisez-moi (rempli à la fin)
readme = wb.create_sheet("Lisez-moi")

# ================================================================== 1. Échelles
ws = sheet("Échelles", "Échelles de mesure",
           "Les trois échelles qui structurent toute l'analyse : gravité d'un incident (1-4), niveau de risque d'un pays (1-5) "
           "et cotation de l'Amirauté (fiabilité A-F × crédibilité 1-6). Relisez les définitions : ce sont elles qui s'affichent "
           "aux utilisateurs et qui guident la notation automatique.",
           "Définitions des échelles de gravité, de risque pays et de cotation")
widths(ws, 10, 18, 60, 60, 36)
r = section(ws, 5, "Gravité d'un incident (1 à 4)")
header(ws, r, ["Valeur", "Libellé", "Repères utilisés aujourd'hui", "Votre définition / correction", "Commentaire"])
SEV_HINT = {1: "Incident isolé, sans victime ou impact local limité ; catégories courantes (crime, troubles, tempête) par défaut.",
            2: "Incident notable : victimes, séisme M5-6, inondation, attaque armée par défaut, bilan 1-9 morts.",
            3: "Incident grave : terrorisme par défaut, bilan ≥ 10 morts, séisme M6-7, alerte PAGER orange, escalade (« massacre », « coup d'État »…).",
            4: "Incident critique : bilan ≥ 50 morts, séisme ≥ M7 ou PAGER rouge, catastrophe majeure, attentat de masse."}
for i in range(1, 5):
    r += 1
    put(ws, r, 1, i, bold=True)
    put(ws, r, 2, SEV_NAMES[i])
    put(ws, r, 3, SEV_HINT[i])
    put(ws, r, 4, None, "in")
    put(ws, r, 5, None, "in")
r = section(ws, r + 2, "Niveau de risque pays (1 à 5)")
header(ws, r, ["Niveau", "Libellé", "Description affichée aujourd'hui", "Votre description", "Commentaire"])
for i in range(1, 6):
    r += 1
    put(ws, r, 1, i, bold=True)
    put(ws, r, 2, LEVEL_NAMES[i])
    put(ws, r, 3, model.RISK_LEVELS[i]["desc_fr"])
    put(ws, r, 4, None, "in")
    put(ws, r, 5, None, "in")
r = section(ws, r + 2, "Cotation de l'Amirauté – fiabilité de la source (lettre)")
header(ws, r, ["Lettre", "Libellé", "Qui reçoit cette lettre aujourd'hui", "Votre règle", "Commentaire"])
for k, lab, who in [("A", "Totalement fiable", "Capteurs et organismes officiels : USGS, GDACS, NASA, OMS, UCDP, NWS, Meteoalarm, CISA, NOAA, ministères (voir onglet Cotation)"),
                    ("B", "Habituellement fiable", "Médias de référence (catalogue de l'onglet Médias, liste « référence » de l'onglet Cotation)"),
                    ("C", "Assez fiable", "Toute autre presse identifiée"),
                    ("D", "Pas toujours fiable", "Détection automatique GDELT (codage machine de la presse mondiale) ; médias notés D dans le catalogue"),
                    ("E", "Peu fiable", "Réseaux sociaux, canaux Telegram"),
                    ("F", "Fiabilité inconnue", "Source non identifiée")]:
    r += 1
    put(ws, r, 1, k, bold=True)
    put(ws, r, 2, lab)
    put(ws, r, 3, who)
    put(ws, r, 4, None, "in")
    put(ws, r, 5, None, "in")
r = section(ws, r + 2, "Cotation de l'Amirauté – crédibilité de l'information (chiffre)")
header(ws, r, ["Chiffre", "Libellé", "Règle automatique actuelle", "Votre règle", "Commentaire"])
for k, lab, rule in [(1, "Confirmée", "Validée par l'analyste, ou mesure officielle directe (capteur)"),
                     (2, "Probablement vraie", "Au moins 3 sources indépendantes, confiance non faible"),
                     (3, "Possiblement vraie", "2 sources, ou confiance haute, ou 1 source cotée A/B"),
                     (4, "Douteuse", "Source unique, confiance faible ou source cotée C à F"),
                     (5, "Improbable", "Infirmée par l'analyste (fausse alerte)"),
                     (6, "Invérifiable", "Non utilisé automatiquement aujourd'hui")]:
    r += 1
    put(ws, r, 1, k, bold=True)
    put(ws, r, 2, lab)
    put(ws, r, 3, rule)
    put(ws, r, 4, None, "in")
    put(ws, r, 5, None, "in")
finish(ws, "A5")

# ================================================================== 2. Catégories
ws = sheet("Catégories", "Catégories d'incidents",
           "Les 20 catégories de la carte. Pour chacune : groupe (qui décide si elle compte dans la note de risque), gravité de "
           "départ quand un titre de presse est classé dans cette catégorie, et ordre de priorité quand un titre correspond à "
           "plusieurs catégories (1 = gagne). Proposez renommages, fusions, scissions ou nouvelles catégories en bas du tableau.",
           "Catégories, groupes, gravité de départ, priorité de classement")
widths(ws, 16, 24, 24, 18, 11, 11, 9, 9, 14, 12, 30, 36)
r = 5
header(ws, r, ["Clé technique", "Libellé FR", "Libellé EN", "Groupe", "Gravité de départ (presse)", "Votre gravité de départ",
               "Priorité actuelle", "Votre priorité", "Compte dans la note pays ?", "Garder la catégorie ?", "Votre libellé / regroupement", "Commentaire"])
prio = {c: i + 1 for i, c in enumerate(press.PRIORITY)}
first_cat = r + 1
for key, c in model.CATEGORIES.items():
    r += 1
    grp = model.GROUPS.get(c["group"], {}).get("fr", c["group"])
    put(ws, r, 1, key)
    put(ws, r, 2, c["fr"], bold=True)
    put(ws, r, 3, c["en"])
    put(ws, r, 4, grp)
    put(ws, r, 5, press.BASE_SEVERITY.get(key, ""))
    put(ws, r, 6, None, "in")
    put(ws, r, 7, prio.get(key, ""))
    put(ws, r, 8, None, "in")
    put(ws, r, 9, "non" if c["group"] == "diplomatic" else "oui")
    put(ws, r, 10, "oui", "in")
    put(ws, r, 11, None, "in")
    put(ws, r, 12, None, "in")
last_cat = r
dv_list(ws, f"F{first_cat}:F{last_cat + 10}", ["1", "2", "3", "4"])
dv_list(ws, f"J{first_cat}:J{last_cat + 10}", ["oui", "non", "fusionner", "scinder"])
r += 1
put(ws, r, 1, "(nouvelle)", "note")
for col in range(2, 13):
    put(ws, r, col, None, "in")
put(ws, r, 12, "Exemple : « Piraterie maritime » – à séparer de Criminalité ?", "in")
for _ in range(5):
    r += 1
    for col in range(1, 13):
        put(ws, r, col, None, "in")
r = section(ws, r + 2, "Groupes de catégories")
header(ws, r, ["Clé", "Groupe", "Effet aujourd'hui", "Votre proposition"])
for key, g in model.GROUPS.items():
    r += 1
    eff = {"security": "Activité sécuritaire (note pays 30 %, Pulse 30 %)", "political": "Activité sécuritaire (comme Sécurité)",
           "natural": "Aléas (note pays 10 %, Pulse 10 %)", "health": "Aléas (comme Catastrophes)",
           "infrastructure": "Aléas (comme Catastrophes)", "diplomatic": "Affiché, mais hors note pays et hors Pulse"}[key]
    put(ws, r, 1, key)
    put(ws, r, 2, g["fr"], bold=True)
    put(ws, r, 3, eff)
    put(ws, r, 4, None, "in")
finish(ws, "C6", f"A5:L{last_cat}")

# ================================================================== 3. Gravité (règles)
ws = sheet("Règles de gravité", "Règles de gravité automatique",
           "Comment l'outil fixe la gravité (1-4) d'un incident selon sa source. Modifiez les seuils dans les colonnes jaunes. "
           "Pour la presse : gravité de départ de la catégorie (onglet Catégories), +1 si un mot d'escalade est présent, puis "
           "relèvement selon le bilan humain cité dans le titre.",
           "Seuils de gravité par source (bilans, magnitudes, alertes officielles)")
widths(ws, 26, 46, 16, 16, 40)
r = 5
header(ws, r, ["Source", "Condition", "Gravité actuelle", "Votre gravité / seuil", "Commentaire"])
rules = [
    ("Presse – bilan", "Titre citant au moins 10 morts", 3), ("Presse – bilan", "Titre citant au moins 50 morts", 4),
    ("Presse – bilan", "Titre citant 1 à 9 morts", 2),
    ("Presse – escalade", "Mot d'escalade (" + ", ".join(press.ESCALATE[:8]) + "…)", "+1"),
    ("USGS (séismes)", "Magnitude < 5", 1), ("USGS (séismes)", "Magnitude 5 à 5,9", 2), ("USGS (séismes)", "Magnitude 6 à 6,9", 3),
    ("USGS (séismes)", "Magnitude ≥ 7", 4), ("USGS (séismes)", "Alerte PAGER verte / jaune / orange / rouge", "1 / 2 / 3 / 4"),
    ("USGS (séismes)", "Alerte tsunami", "+1"), ("Presse (séismes)", "Magnitude < 5 citée par la presse : écartée (déjà couverte par l'USGS)", "écarté"),
    ("GDACS (catastrophes)", "Alerte verte / orange / rouge", "1 / 3 / 4"),
    ("NWS / Meteoalarm (météo)", "Vigilance orange (Severe) / rouge (Extreme)", "3 / 4"),
    ("OMS (santé)", "Épidémie signalée (bulletin officiel)", 2),
    ("OMS (santé)", "Maladie à haut risque (Ebola, Marburg, peste, Nipah, MERS, grippe aviaire, Lassa…)", 3),
    ("GDELT (détection automatique)", "Gravité de base selon le type d'événement codé (CAMEO) ; hors zone de conflit, « usage de la force » devient Attaque (gravité ≤ 3)", "1 à 3"),
    ("GDELT (détection automatique)", "Au moins 20 articles et 3 médias différents", "+1"),
    ("Avis aux voyageurs", "Ne modifient pas la gravité d'un incident : ils pèsent dans la note pays", "—"),
]
for src, cond, g in rules:
    r += 1
    put(ws, r, 1, src, bold=True)
    put(ws, r, 2, cond)
    put(ws, r, 3, g)
    put(ws, r, 4, None, "in")
    put(ws, r, 5, None, "in")
r += 1
for col in range(1, 6):
    put(ws, r, col, None, "in")
put(ws, r, 1, "(nouvelle règle)", "in")
finish(ws, "A6")

# ================================================================== 4. Note pays (simulateur)
ws = sheet("Note pays", "Note de risque pays (1 à 5) – réglages et simulateur",
           "La note combine trois composantes entre 0 et 1 : avis officiels (moyenne des ministères), activité sécuritaire et "
           "aléas des derniers jours (échelle logarithmique : quelques incidents graves pèsent plus que beaucoup de mineurs). "
           "Changez les réglages jaunes : le simulateur en bas recalcule les exemples. Ajoutez vos propres cas pour tester.",
           "Poids, seuils et plancher de la note pays ; niveaux imposés par l'analyste")
widths(ws, 30, 14, 14, 52, 9, 9, 9, 9, 9, 9, 9, 9, 9, 9, 11, 11, 11, 11, 11, 11, 10, 11, 9)
r = section(ws, 5, "Réglages")
header(ws, r, ["Paramètre", "Valeur actuelle", "Valeur testée", "Explication"])
P = {}
params = [
    ("w_adv", "Poids des avis officiels", RISK["weights"]["advisories"], "Part de la note qui vient des avis des ministères (France, Royaume-Uni, États-Unis, Allemagne, Canada)."),
    ("w_sec", "Poids de l'activité sécuritaire", RISK["weights"]["security"], "Incidents des groupes Sécurité et Politique sur la fenêtre."),
    ("w_haz", "Poids des aléas", RISK["weights"]["hazards"], "Catastrophes naturelles, santé, infrastructures sur la fenêtre."),
    ("window", "Fenêtre (jours)", RISK["window_days"], "Période d'incidents prise en compte. Information (le simulateur saisit directement les incidents de la fenêtre)."),
    ("p1", "Points d'un incident de gravité 1", RISK["severity_points"]["1"], "Poids de chaque incident selon sa gravité."),
    ("p2", "Points d'un incident de gravité 2", RISK["severity_points"]["2"], ""),
    ("p3", "Points d'un incident de gravité 3", RISK["severity_points"]["3"], ""),
    ("p4", "Points d'un incident de gravité 4", RISK["severity_points"]["4"], ""),
    ("sat", "Saturation (points pour une activité maximale)", RISK["saturation_points"], "À partir de ce total de points, la composante vaut 1 (100 %)."),
    ("t2", "Seuil du niveau 2 (Faible)", RISK["thresholds"][0], "Score minimal (0 à 1) pour atteindre ce niveau."),
    ("t3", "Seuil du niveau 3 (Modéré)", RISK["thresholds"][1], ""),
    ("t4", "Seuil du niveau 4 (Élevé)", RISK["thresholds"][2], ""),
    ("t5", "Seuil du niveau 5 (Extrême)", RISK["thresholds"][3], ""),
    ("auto", "Poids d'une détection automatique non recoupée", RISK["auto_detected_weight"], "Une détection GDELT compte pour cette fraction d'un incident confirmé."),
    ("floor", "Plancher si un ministère dit « ne pas se rendre » (4/4)", 4, "Niveau minimal du pays dans ce cas."),
]
for key, lab, val, expl in params:
    r += 1
    put(ws, r, 1, lab, bold=True)
    put(ws, r, 2, val)
    put(ws, r, 3, val, "in")
    put(ws, r, 4, expl, "note")
    P[key] = f"$C${r}"
r += 1
put(ws, r, 1, "Contrôle : somme des poids", bold=True)
put(ws, r, 3, f"={P['w_adv']}+{P['w_sec']}+{P['w_haz']}", "calc", "0.00")
put(ws, r, 4, f'=IF(ABS(C{r}-1)<0.001,"OK : les poids font 100 %","Attention : les poids doivent faire 1 (100 %)")', "calc")
r = section(ws, r + 2, "Simulateur (exemples fictifs – remplacez ou ajoutez vos propres cas)")
r_head = r
header(ws, r, ["Cas", "Avis US (1-4)", "Avis FCDO (1-4)", "Avis MEAE (1-4)", "Avis DE (1-4)", "Sécu. g1", "Sécu. g2", "Sécu. g3", "Sécu. g4",
               "Part auto (0-1)", "Aléas g1", "Aléas g2", "Aléas g3", "Aléas g4",
               "Avis (0-1)", "Points sécu.", "Activité sécu. (0-1)", "Points aléas", "Activité aléas (0-1)", "Score (0-1)",
               "Niveau calculé", "Votre niveau attendu", "Écart"])
cases = [("Pays en guerre (type Mali)", 4, 4, 4, 4, 3, 8, 6, 2, 0.3, 0, 1, 0, 0),
         ("Pays sous tension (type Kenya)", 2, 1, 2, 2, 4, 5, 2, 0, 0.4, 1, 1, 0, 0),
         ("Pays stable, manifestations (type France)", 1, 1, 1, 1, 3, 4, 0, 0, 0.5, 2, 1, 0, 0),
         ("Pays calme, séisme majeur", 1, 1, 1, 1, 0, 0, 0, 0, 0, 0, 2, 1, 1),
         ("Pays sans avis officiel, violences", None, None, None, None, 2, 6, 4, 1, 0.6, 0, 0, 0, 0),
         ("Votre cas", None, None, None, None, None, None, None, None, None, None, None, None, None)]
first_case = r + 1
for case in cases:
    r += 1
    for i, v in enumerate(case):
        put(ws, r, i + 1, v, "in")
    B, C_, D_, E, Fc, G, H, I_, J, K, L, M, N = [f"{get_column_letter(c)}{r}" for c in range(2, 15)]
    put(ws, r, 15, f'=IF(COUNT(B{r}:E{r})=0,"",(SUM(B{r}:E{r})-COUNT(B{r}:E{r}))/(3*COUNT(B{r}:E{r})))', "calc", "0.00")
    put(ws, r, 16, f"=({Fc}*{P['p1']}+{G}*{P['p2']}+{H}*{P['p3']}+{I_}*{P['p4']})*(1-{J}+{J}*{P['auto']})", "calc", "0.0")
    put(ws, r, 17, f"=IF(P{r}<=0,0,MIN(1,LN(1+P{r})/LN(1+{P['sat']})))", "calc", "0.00")
    put(ws, r, 18, f"={K}*{P['p1']}+{L}*{P['p2']}+{M}*{P['p3']}+{N}*{P['p4']}", "calc", "0.0")
    put(ws, r, 19, f"=IF(R{r}<=0,0,MIN(1,LN(1+R{r})/LN(1+{P['sat']})))", "calc", "0.00")
    put(ws, r, 20, f'=IF(O{r}="",({P["w_sec"]}*Q{r}+{P["w_haz"]}*S{r})/({P["w_sec"]}+{P["w_haz"]}),{P["w_adv"]}*O{r}+{P["w_sec"]}*Q{r}+{P["w_haz"]}*S{r})', "calc", "0.00")
    put(ws, r, 21, f'=MAX(1+(T{r}>={P["t2"]})+(T{r}>={P["t3"]})+(T{r}>={P["t4"]})+(T{r}>={P["t5"]}),IF(MAX(B{r}:E{r})>=4,{P["floor"]},1))', "calc", "0", bold=True)
    put(ws, r, 22, None, "in")
    put(ws, r, 23, f'=IF(V{r}="","",U{r}-V{r})', "calc", "+0;-0;0")
last_case = r
ws.conditional_formatting.add(f"W{first_case}:W{last_case}", CellIsRule(operator="notEqual", formula=["0"], fill=PatternFill("solid", fgColor="F8D7DA")))
put(ws, r + 1, 1, "Hypothèse du simulateur : tous les incidents saisis sont de confiance « haute » ; seule la part de détections "
    "automatiques est réduite. Avis : laisser vide si le ministère n'a pas d'avis.", "note")
ws.merge_cells(start_row=r + 1, start_column=1, end_row=r + 1, end_column=14)
r = section(ws, r + 3, "Niveaux imposés par l'analyste (remplacent le calcul)")
header(ws, r, ["Pays (code ISO, ex. ML)", "Niveau imposé (1-5)", "Valable jusqu'au", "Justification"])
for iso, lvl in (RISK.get("overrides") or {}).items():
    r += 1
    put(ws, r, 1, iso, "in")
    put(ws, r, 2, lvl, "in")
    put(ws, r, 3, None, "in")
    put(ws, r, 4, None, "in")
r += 1
for col, v in enumerate(["HT", 5, "31/12/2026", "Exemple : capitale contrôlée par les gangs, avis ministériels en retard sur la situation"], 1):
    put(ws, r, col, v, "in")
for _ in range(8):
    r += 1
    for col in range(1, 5):
        put(ws, r, col, None, "in")
finish(ws, "B6")

# ================================================================== 5. Matrice des menaces (rapport pays)
ws = sheet("Matrice menaces", "Matrice des menaces du rapport pays",
           "Dans le rapport pays, chaque menace reçoit un niveau 1-5 : niveau tiré du nombre d'incidents sur 90 jours, relevé par des "
           "indicateurs structurels (avis, groupes armés, stabilité, homicides, motifs cités par le Département d'État). "
           "Le niveau final est le plus élevé des deux. Ajustez seuils et planchers.",
           "Seuils d'incidents, indicateurs structurels et planchers par menace")
widths(ws, 34, 44, 14, 14, 50, 30)
r = section(ws, 5, "Niveau selon le nombre d'incidents sur 90 jours")
header(ws, r, ["Niveau", "Libellé", "Incidents 90 j (minimum) – actuel", "Votre seuil", "Commentaire"])
TH = {}
for lvl, mn in [(2, 1), (3, 5), (4, 20), (5, 60)]:
    r += 1
    put(ws, r, 1, lvl, bold=True)
    put(ws, r, 2, LEVEL_NAMES[lvl])
    put(ws, r, 3, mn)
    put(ws, r, 4, mn, "in")
    put(ws, r, 5, None, "in")
    TH[lvl] = f"$D${r}"
r = section(ws, r + 2, "Stabilité politique (Banque mondiale, 0 = pire, 100 = meilleur) → plancher « Instabilité politique »")
header(ws, r, ["Plancher", "Libellé", "Si stabilité inférieure à – actuel", "Votre seuil", "Commentaire"])
ST = {}
for lvl, th in [(5, 15), (4, 30), (3, 50), (2, 70)]:
    r += 1
    put(ws, r, 1, lvl, bold=True)
    put(ws, r, 2, LEVEL_NAMES[lvl])
    put(ws, r, 3, th)
    put(ws, r, 4, th, "in")
    put(ws, r, 5, None, "in")
    ST[lvl] = f"$D${r}"
r = section(ws, r + 2, "Homicides pour 100 000 habitants → plancher « Criminalité »")
header(ws, r, ["Plancher", "Libellé", "Si homicides au moins – actuel", "Votre seuil", "Commentaire"])
HO = {}
for lvl, th in [(5, 20), (4, 10), (3, 5), (2, 2)]:
    r += 1
    put(ws, r, 1, lvl, bold=True)
    put(ws, r, 2, LEVEL_NAMES[lvl])
    put(ws, r, 3, th)
    put(ws, r, 4, th, "in")
    put(ws, r, 5, None, "in")
    HO[lvl] = f"$D${r}"
r = section(ws, r + 2, "Planchers structurels par menace")
header(ws, r, ["Menace", "Indicateur déclencheur", "Plancher actuel", "Votre plancher", "Commentaire", "Catégories d'incidents comptées"])
floors = [
    ("Terrorisme", "Groupes terroristes recensés (CIA Factbook)", "3 (4 si un avis 4/4)", "terrorism"),
    ("Terrorisme", "Terrorisme cité par le Département d'État", 3, ""),
    ("Conflit armé / violence politique", "Au moins un ministère : « ne pas se rendre » (4/4)", 4, "armed_conflict"),
    ("Attaques, enlèvements, violences armées", "Enlèvements cités par le Département d'État", "4 (5 si un avis 4/4)", "attack"),
    ("Troubles civils, manifestations, grèves", "Troubles civils cités par le Département d'État", 3, "unrest"),
    ("Instabilité politique", "Stabilité politique (tableau ci-dessus)", "selon seuils", "political"),
    ("Criminalité", "Homicides (tableau ci-dessus)", "selon seuils", "crime"),
    ("Criminalité", "Criminalité citée par le Département d'État", 3, ""),
    ("Cyber et infrastructures", "Menace mondiale de fond (tous pays)", 2, "cyber, infrastructure"),
    ("Catastrophes naturelles et climat", "Aléas connus (CIA Factbook)", 2, "earthquake, cyclone, storm, flood, wildfire, volcano, drought, landslide, extreme_temp"),
    ("Catastrophes naturelles et climat", "Catastrophes naturelles citées par le Département d'État", 3, ""),
    ("Santé et épidémies", "Offre de soins limitée citée par le Département d'État", 3, "health"),
    ("Détention arbitraire (fiche Sûreté)", "Détention arbitraire citée par le Département d'État", 4, "—"),
]
for m, ind, fl, cats in floors:
    r += 1
    put(ws, r, 1, m, bold=True)
    put(ws, r, 2, ind)
    put(ws, r, 3, fl)
    put(ws, r, 4, None, "in")
    put(ws, r, 5, None, "in")
    put(ws, r, 6, cats)
r += 1
for col in range(1, 7):
    put(ws, r, col, None, "in")
put(ws, r, 1, "(nouvelle menace ou règle)", "in")
r = section(ws, r + 2, "Simulateur")
header(ws, r, ["Cas", "Incidents 90 j", "Niveau incidents", "Plancher structurel (saisi)", "Niveau final", "Stabilité politique (0-100)"])
for lab, n, fl in [("Terrorisme – pays du Sahel", 35, 4), ("Criminalité – grande ville d'Amérique latine", 12, 4), ("Troubles – pays européen", 6, 1), ("Votre cas", None, None)]:
    r += 1
    put(ws, r, 1, lab, "in")
    put(ws, r, 2, n, "in")
    put(ws, r, 3, f"=1+(B{r}>={TH[2]})+(B{r}>={TH[3]})+(B{r}>={TH[4]})+(B{r}>={TH[5]})", "calc", "0")
    put(ws, r, 4, fl, "in")
    put(ws, r, 5, f'=MAX(C{r},IF(D{r}="",1,D{r}))', "calc", "0", bold=True)
r += 1
put(ws, r, 1, "Calculette stabilité → plancher : saisir une stabilité", "note")
put(ws, r, 2, 22, "in")
put(ws, r, 3, f"=IF(B{r}<{ST[5]},5,IF(B{r}<{ST[4]},4,IF(B{r}<{ST[3]},3,IF(B{r}<{ST[2]},2,1))))", "calc", "0", bold=True)
r += 1
put(ws, r, 1, "Calculette homicides → plancher : saisir un taux", "note")
put(ws, r, 2, 7.5, "in")
put(ws, r, 3, f"=IF(B{r}>={HO[5]},5,IF(B{r}>={HO[4]},4,IF(B{r}>={HO[3]},3,IF(B{r}>={HO[2]},2,1))))", "calc", "0", bold=True)
finish(ws, "A5")

# ================================================================== 6. Pulse
ws = sheet("Pulse", "Indice de stabilité Pulse (0 à 100)",
           "Le Pulse suit la dynamique d'un pays : 100 = stable, 0 = très instable. Instabilité = somme pondérée de 4 composantes ; "
           "Pulse = 100 × (1 − instabilité). L'anomalie compare les 7 derniers jours à la moyenne des 23 jours précédents.",
           "Poids du Pulse, calcul de l'anomalie, seuils d'alerte")
widths(ws, 34, 14, 14, 50, 12, 12, 12, 12, 12, 12, 12, 12, 12, 12)
r = section(ws, 5, "Réglages")
header(ws, r, ["Paramètre", "Valeur actuelle", "Valeur testée", "Explication"])
Q = {}
pulse_cfg = SETTINGS.get("pulse") or {}
for key, lab, val, expl in [
        ("w_adv", "Poids des avis officiels", 0.45, "Composante avis de la note pays (0-1)."),
        ("w_sec", "Poids de l'activité sécuritaire 7 j", 0.30, "Composante sécurité de la note pays (0-1)."),
        ("w_ano", "Poids de l'anomalie", 0.15, "Hausse brutale par rapport à la normale du pays."),
        ("w_haz", "Poids des catastrophes et santé 7 j", 0.10, "Composante aléas de la note pays (0-1)."),
        ("min_n", "Points minimaux pour parler d'anomalie", 3, "En dessous (7 j), anomalie = 0."),
        ("low_base", "Normale « quasi nulle » (points / 7 j)", 0.5, "Si la normale est inférieure, anomalie = points récents / diviseur ci-dessous."),
        ("low_div", "Diviseur si normale quasi nulle", 20, "20 points en 7 j sans antécédent = anomalie maximale."),
        ("max_ratio", "Multiple de la normale = anomalie maximale", 3, "3 = trois fois plus d'incidents que d'habitude → anomalie 100 %."),
        ("alert_th", "Alerte si le Pulse passe sous", pulse_cfg.get("threshold", 40), "Alerte Telegram (pays suivis)."),
        ("alert_drop", "Alerte si chute en 7 jours d'au moins", pulse_cfg.get("drop_alert", 12), "Points de Pulse perdus en 7 jours."),
        ("gap", "Délai minimal entre deux alertes (jours)", pulse_cfg.get("min_gap_days", 5), "")]:
    r += 1
    put(ws, r, 1, lab, bold=True)
    put(ws, r, 2, val)
    put(ws, r, 3, val, "in")
    put(ws, r, 4, expl, "note")
    Q[key] = f"$C${r}"
r += 1
put(ws, r, 1, "Contrôle : somme des poids", bold=True)
put(ws, r, 3, f"={Q['w_adv']}+{Q['w_sec']}+{Q['w_ano']}+{Q['w_haz']}", "calc", "0.00")
put(ws, r, 4, f'=IF(ABS(C{r}-1)<0.001,"OK","Attention : les poids doivent faire 1")', "calc")
r = section(ws, r + 2, "Simulateur")
header(ws, r, ["Cas", "Avis (0-1)", "Sécurité 7 j (0-1)", "Aléas 7 j (0-1)", "Points 7 derniers jours", "Points 23 jours précédents",
               "Normale sur 7 j", "Anomalie (0-1)", "Instabilité", "Pulse", "Alerte (seuil) ?", "Votre Pulse attendu"])
for lab, a, s_, h, rec, pri in [("Pays stable", 0, 0.1, 0.05, 2, 10), ("Flambée soudaine de violence", 0.33, 0.6, 0, 60, 30),
                                ("Pays en guerre installée", 1, 0.9, 0.1, 120, 380), ("Catastrophe majeure", 0, 0.1, 0.9, 1, 3), ("Votre cas", None, None, None, None, None)]:
    r += 1
    for i, v in enumerate([lab, a, s_, h, rec, pri]):
        put(ws, r, i + 1, v, "in")
    put(ws, r, 7, f"=F{r}/23*7", "calc", "0.0")
    put(ws, r, 8, f"=IF(E{r}<{Q['min_n']},0,IF(G{r}<={Q['low_base']},MIN(1,E{r}/{Q['low_div']}),MAX(0,MIN(1,(E{r}/G{r}-1)/({Q['max_ratio']}-1)))))", "calc", "0.00")
    put(ws, r, 9, f"={Q['w_adv']}*B{r}+{Q['w_sec']}*C{r}+{Q['w_ano']}*H{r}+{Q['w_haz']}*D{r}", "calc", "0.00")
    put(ws, r, 10, f"=ROUND(100*(1-MIN(1,I{r})),0)", "calc", "0", bold=True)
    put(ws, r, 11, f'=IF(J{r}<{Q["alert_th"]},"oui","non")', "calc")
    put(ws, r, 12, None, "in")
finish(ws, "A5")

# ================================================================== 7. Cotation
ws = sheet("Cotation", "Cotation de l'Amirauté : règles et apprentissage",
           "Qui est classé A (officiel) ou B (référence), et comment l'outil ajuste la lettre d'une source à partir de vos décisions "
           "« Valider / Fausse alerte ». La lettre de chaque média du catalogue se règle dans l'onglet Médias.",
           "Sources officielles (A), médias de référence (B), règles d'apprentissage")
widths(ws, 34, 16, 16, 50)
r = section(ws, 5, "Apprentissage à partir de vos validations")
header(ws, r, ["Règle", "Valeur actuelle", "Votre valeur", "Effet"])
for lab, val, eff in [("Décisions minimales avant ajustement", 5, "En dessous, la lettre de la source ne bouge pas."),
                      ("Taux de fausses alertes → perd 1 lettre", 0.5, "Ex. B devient C."),
                      ("Taux de fausses alertes → perd 2 lettres", 0.8, "Ex. B devient D."),
                      ("Décisions minimales pour gagner une lettre", 8, ""),
                      ("Taux de confirmation → gagne 1 lettre", 0.9, "Jamais au-delà de B (A reste réservé aux sources officielles).")]:
    r += 1
    put(ws, r, 1, lab, bold=True)
    put(ws, r, 2, val, fmt="0%" if val < 1 else "0")
    put(ws, r, 3, val, "in", "0%" if val < 1 else "0")
    put(ws, r, 4, eff, "note")
r = section(ws, r + 2, "Sources cotées A (officielles)")
header(ws, r, ["Source", "Mesure directe (crédibilité 1 d'office) ?", "Garder en A ?", "Commentaire"])
for s_ in sorted(quality.OFFICIAL):
    r += 1
    put(ws, r, 1, s_, bold=True)
    put(ws, r, 2, "oui" if s_ in quality.SENSORS else "non")
    put(ws, r, 3, "oui", "in")
    put(ws, r, 4, None, "in")
for _ in range(3):
    r += 1
    for col in range(1, 5):
        put(ws, r, col, None, "in")
r = section(ws, r + 2, "Médias de référence hors catalogue (cotés B)")
header(ws, r, ["Média", "Lettre actuelle", "Votre lettre", "Commentaire"])
for s_ in sorted(quality.REFERENCE_EXTRA):
    r += 1
    put(ws, r, 1, s_, bold=True)
    put(ws, r, 2, "B")
    put(ws, r, 3, None, "in")
    put(ws, r, 4, None, "in")
dv_list(ws, f"C6:C{r + 5}", list("ABCDEF") + ["oui", "non"])
finish(ws, "A5")

# ================================================================== 8. Médias
ws = sheet("Médias", "Catalogue des médias par pays",
           f"{sum(len(v) for v in OUTLETS.values())} médias de référence dans {len(OUTLETS)} pays, interrogés via Google News. Lettre actuelle : "
           "« B » par défaut, sauf mention. Notez D les médias d'État ou partisans, C les médias moins établis, et signalez les "
           "médias à retirer ou à ajouter (lignes vides en bas). Utilisez les filtres de la ligne d'en-tête.",
           "Fiabilité (lettre) de chaque média, médias à retirer ou à ajouter")
widths(ws, 8, 22, 34, 28, 9, 10, 12, 12, 40)
r = 5
header(ws, r, ["Pays", "Nom du pays", "Média", "Domaine", "Langue", "Lettre actuelle", "Votre lettre", "Garder ?", "Commentaire"])
for iso in sorted(OUTLETS, key=lambda i: CNAME.get(i, i)):
    for o in OUTLETS[iso]:
        r += 1
        letter = o[4] if len(o) > 4 and o[4] in "ABCDEF" else "B"
        put(ws, r, 1, iso)
        put(ws, r, 2, CNAME.get(iso, iso))
        put(ws, r, 3, o[0], bold=True)
        put(ws, r, 4, o[1])
        put(ws, r, 5, o[2] if len(o) > 2 else "")
        put(ws, r, 6, letter)
        put(ws, r, 7, None, "in")
        put(ws, r, 8, None, "in")
        put(ws, r, 9, None, "in")
last = r
for _ in range(20):
    r += 1
    for col in range(1, 10):
        put(ws, r, col, None, "in")
dv_list(ws, f"G6:G{r}", list("ABCDEF"))
dv_list(ws, f"H6:H{r}", ["oui", "non", "ajouter"])
finish(ws, "D6", f"A5:I{last}")

# ================================================================== 9. Sources
ws = sheet("Sources", "Sources de données",
           f"{len(SOURCES)} sources déclarées ({sum(1 for s_ in SOURCES if s_.get('enabled'))} actives). Pour chacune : votre avis sur sa fiabilité, "
           "sa pertinence pour vos clients, et la licence (usage commercial possible ?). Proposez de nouvelles sources en bas.",
           "Fiabilité, pertinence, licence et maintien de chaque source")
widths(ws, 26, 34, 14, 9, 36, 12, 12, 16, 12, 40)
r = 5
header(ws, r, ["Identifiant", "Nom", "Type", "Active", "Licence déclarée", "Votre fiabilité (A-F)", "Pertinence (1-3)", "Usage commercial OK ?", "Garder ?", "Commentaire"])
for s_ in sorted(SOURCES, key=lambda x: (x["type"], x.get("name", x["id"]))):
    r += 1
    put(ws, r, 1, s_["id"])
    put(ws, r, 2, s_.get("name", s_["id"]), bold=True)
    put(ws, r, 3, s_["type"])
    put(ws, r, 4, "oui" if s_.get("enabled") else "non")
    put(ws, r, 5, s_.get("license", ""))
    for col in range(6, 11):
        put(ws, r, col, None, "in")
last = r
for _ in range(15):
    r += 1
    for col in range(1, 11):
        put(ws, r, col, None, "in")
dv_list(ws, f"F6:F{r}", list("ABCDEF"))
dv_list(ws, f"G6:G{r}", ["1", "2", "3"], "1 = faible, 3 = indispensable")
dv_list(ws, f"H6:H{r}", ["oui", "non", "à vérifier"])
dv_list(ws, f"I6:I{r}", ["oui", "non", "ajouter"])
finish(ws, "C6", f"A5:J{last}")

# ================================================================== 10. Mots-clés
ws = sheet("Mots-clés", "Mots-clés du tri de la presse",
           "Le tri des titres repose sur des listes de racines de mots, toutes langues confondues (un mot court de 4 lettres ou moins "
           "doit être exact ; « $ » à la fin force le mot exact). Listes « catégorie » : un titre qui contient le mot entre dans la "
           "catégorie. Listes « filtre » : elles écartent ou retiennent un titre. Action : garder, supprimer, ou ajouter un mot (lignes vides en bas de liste).",
           "Mots-clés de chaque catégorie et des filtres (faits divers, judiciaire, sport…)")
widths(ws, 26, 46, 30, 12, 30, 40)
FILTER_ROLE = {
    "NOISE_WORDS": "Écarte : sport, culture, exercices, commémorations…", "PRIVATE_WORDS": "Écarte : drames privés, faits divers familiaux",
    "JUDICIAL_WORDS": "Écarte : procès, condamnations, enquêtes (sauf mobilisation)", "RETRO_WORDS": "Écarte : rétrospectives, anniversaires de catastrophes",
    "ECONOMY_WORDS": "Classe en veille économique", "ANIMAL_WORDS": "Écarte : attaques d'animaux", "FIREWORK_WORDS": "Écarte : feux d'artifice, pétards",
    "ACCIDENT_WORDS": "Explosion accidentelle → Infrastructure", "EXPLOSION_WORDS": "Repère une explosion", "INTENT_WORDS": "Explosion intentionnelle → Attaque",
    "VIOLENCE_WORDS": "Signale de la violence (garde une attaque)", "MASS_WORDS": "Signale un nombre important de victimes",
    "PUBLIC_TARGET_WORDS": "Lieu public ou cible visée (garde une attaque)", "SECURITY_TARGET_WORDS": "Cible sécuritaire (police, armée…)",
    "ARMED_GROUP_WORDS": "Groupe armé cité (garde une attaque)", "AFTERMATH_WORDS": "Suites d'un incident (hommage, enquête…) : écarte",
    "THREAT_WORDS": "Simple menace : écarte", "EVACUATION_WORDS": "Évacuation", "PUBLIC_CRIME_WORDS": "Crime touchant l'ordre public (garde)",
    "EXTRA_WORDS": "Compléments", "_NUM_WORDS": "Nombres écrits en toutes lettres (bilans)"}
r = 5
header(ws, r, ["Liste", "Rôle", "Mot-clé (racine)", "Action", "Nouveau mot / correction", "Commentaire"])
lists = [(f"Catégorie : {model.CATEGORIES.get(c, {}).get('fr', c)}", f"Entre dans « {model.CATEGORIES.get(c, {}).get('fr', c)} »", ws_)
         for c, ws_ in press.CATEGORY_WORDS.items()]
lists.append(("Escalade (+1 gravité)", "Relève la gravité d'un cran", press.ESCALATE))
for name in sorted(n for n in dir(press) if n.isupper() and n.endswith("WORDS") and n != "CATEGORY_WORDS"):
    lists.append((f"Filtre : {name.replace('_WORDS', '').strip('_').lower()}", FILTER_ROLE.get(name, ""), getattr(press, name)))
for lname, role, words in lists:
    for w in words:
        r += 1
        put(ws, r, 1, lname)
        put(ws, r, 2, role)
        put(ws, r, 3, w, bold=True)
        put(ws, r, 4, "garder", "in")
        put(ws, r, 5, None, "in")
        put(ws, r, 6, None, "in")
    for _ in range(2):
        r += 1
        put(ws, r, 1, lname, "note")
        put(ws, r, 2, "(ligne libre pour ajouter un mot)", "note")
        put(ws, r, 3, None, "in")
        put(ws, r, 4, None, "in")
        put(ws, r, 5, None, "in")
        put(ws, r, 6, None, "in")
dv_list(ws, f"D6:D{r}", ["garder", "supprimer", "ajouter", "déplacer"])
ws.conditional_formatting.add(f"A6:F{r}", FormulaRule(formula=['$D6="supprimer"'], fill=PatternFill("solid", fgColor="F8D7DA")))
ws.conditional_formatting.add(f"A6:F{r}", FormulaRule(formula=['$D6="ajouter"'], fill=PatternFill("solid", fgColor="D4EDDA")))
finish(ws, "D6", f"A5:F{r}")

# ================================================================== 11. Jeu de test
ws = sheet("Jeu de test", "Jeu de test du tri (titres réels classés à la main)",
           "120 titres réels classés par Claude : faut-il les garder (incident utile à une organisation ou un voyageur) et dans quelle "
           "catégorie ? Donnez votre avis dans les colonnes jaunes : les désaccords s'affichent en rouge et serviront à corriger le tri "
           "et à mesurer sa fiabilité. Ajoutez des titres que vous avez vus passer (bons ou mauvais) en bas.",
           "Validation des 120 titres de référence ; ajout de nouveaux cas")
widths(ws, 70, 7, 11, 16, 12, 16, 12, 36)
put(ws, 5, 1, "Titres revus", bold=True)
put(ws, 6, 1, "Désaccords sur « garder »", bold=True)
put(ws, 7, 1, "Désaccords sur la catégorie", bold=True)
r = 9
header(ws, r, ["Titre", "Pays", "Garder ? (Claude)", "Catégorie (Claude)", "Garder ? (vous)", "Catégorie (vous)", "Désaccord", "Commentaire"])
first = r + 1
cats = list(model.CATEGORIES)
for g in GOLD:
    r += 1
    put(ws, r, 1, g["titre"])
    put(ws, r, 2, g["pays"])
    put(ws, r, 3, "oui" if g["garder"].strip() == "1" else "non")
    put(ws, r, 4, g["categorie"])
    put(ws, r, 5, None, "in")
    put(ws, r, 6, None, "in")
    put(ws, r, 7, f'=IF(E{r}="","",IF(E{r}<>C{r},"garder",IF(AND(E{r}="oui",F{r}<>"",F{r}<>D{r}),"catégorie","")))', "calc")
    put(ws, r, 8, g.get("commentaire") or None, "in")
for _ in range(40):
    r += 1
    for col in (1, 2, 3, 4, 5, 6, 8):
        put(ws, r, col, None, "in")
    put(ws, r, 7, f'=IF(E{r}="","",IF(E{r}<>C{r},"garder",IF(AND(E{r}="oui",F{r}<>"",F{r}<>D{r}),"catégorie","")))', "calc")
put(ws, 5, 2, f'=COUNTA(E{first}:E{r})', "calc", "0")
put(ws, 6, 2, f'=COUNTIF(G{first}:G{r},"garder")', "calc", "0")
put(ws, 7, 2, f'=COUNTIF(G{first}:G{r},"catégorie")', "calc", "0")
dv_list(ws, f"C{first}:C{r}", ["oui", "non"])
dv_list(ws, f"E{first}:E{r}", ["oui", "non"])
dvc = DataValidation(type="list", formula1=f"=Catégories!$A$6:$A${5 + len(cats)}", allow_blank=True)
ws.add_data_validation(dvc)
dvc.add(f"F{first}:F{r}")
dvc.add(f"D{first + len(GOLD)}:D{r}")
ws.conditional_formatting.add(f"G{first}:G{r}", CellIsRule(operator="notEqual", formula=['""'], fill=PatternFill("solid", fgColor="F8D7DA")))
finish(ws, "B10", f"A9:H{first + len(GOLD) - 1}")

# ================================================================== 12. Fiabilité et alertes
ws = sheet("Fiabilité & alertes", "Mode fiable, incidents « en cours », alertes, dédoublonnage",
           "Règles qui décident ce que l'utilisateur voit par défaut, ce qui est présenté comme « en cours », ce qui déclenche une "
           "alerte Telegram / e-mail, et quand deux signalements sont fusionnés en un seul incident.",
           "Règles d'affichage, d'alerte et de fusion des doublons")
widths(ws, 40, 50, 18, 18, 36)
r = section(ws, 5, "Mode fiable (activé par défaut)")
header(ws, r, ["Règle", "Détail actuel", "Valeur actuelle", "Votre valeur", "Commentaire"])
rules = [("Détections masquées", "Détections automatiques cotées D, E ou F", "D, E, F"),
         ("… sauf si recoupée", "Signalée par au moins 2 médias différents", "oui"),
         ("… sauf si confiance haute", "Confiance « haute »", "oui"),
         ("… sauf si proche d'un site", "Dans le rayon de vigilance d'un site ou d'un trajet", "oui"),
         ("… sauf si validée", "Validée par l'analyste", "oui")]
for a, b, c in rules:
    r += 1
    put(ws, r, 1, a, bold=True)
    put(ws, r, 2, b)
    put(ws, r, 3, c)
    put(ws, r, 4, None, "in")
    put(ws, r, 5, None, "in")
r = section(ws, r + 2, "Onglet « En cours » (incidents de moins de 72 h retenus si au moins une condition)")
header(ws, r, ["Condition", "Détail actuel", "Valeur actuelle", "Votre valeur", "Commentaire"])
for a, b, c in [("Fenêtre", "Âge maximal de l'incident", "72 h"), ("Critique", "Gravité", "4"),
                ("Élevé et fiable", "Gravité minimale, confiance non faible", "3"), ("Catastrophe GDACS", "Gravité minimale", "3"),
                ("Événement NASA EONET actif", "Gravité minimale", "2"), ("Recoupé (plusieurs sources)", "Gravité minimale", "2"),
                ("Évolutif", "Incident suivi depuis plus de N heures, gravité ≥ 2, fiable", "6 h"),
                ("Proche de vos sites", "Gravité minimale, fiable", "2"), ("Crise majeure", "Au moins 1 incident critique ou N incidents", "5"),
                ("Crise", "Nombre d'incidents liés", "2")]:
    r += 1
    put(ws, r, 1, a, bold=True)
    put(ws, r, 2, b)
    put(ws, r, 3, c)
    put(ws, r, 4, None, "in")
    put(ws, r, 5, None, "in")
alerts = SETTINGS.get("alerts") or {}
r = section(ws, r + 2, "Alertes Telegram et e-mail")
header(ws, r, ["Réglage", "Détail", "Valeur actuelle", "Votre valeur", "Commentaire"])
for a, b, c in [("Gravité minimale d'un incident (monde)", "Tout nouvel incident au moins de cette gravité", alerts.get("min_severity", 3)),
                ("Âge maximal", "Heures depuis l'incident", alerts.get("max_age_hours", 6)),
                ("Pays", "Vide = tous les pays", ", ".join(alerts.get("countries") or []) or "tous"),
                ("Catégories", "Vide = toutes", ", ".join(alerts.get("categories") or []) or "toutes"),
                ("Inclure les détections automatiques", "", "oui" if alerts.get("include_auto") else "non"),
                ("Gravité minimale près d'un site", "Alertes de proximité de vos sites", SETTINGS.get("site_alert_min_severity", 2)),
                ("Rayon de vigilance par défaut d'un site (km)", "Modifiable site par site", 50)]:
    r += 1
    put(ws, r, 1, a, bold=True)
    put(ws, r, 2, b)
    put(ws, r, 3, c)
    put(ws, r, 4, None, "in")
    put(ws, r, 5, None, "in")
r = section(ws, r + 2, "Fusion des doublons : deux signalements = un seul incident s'ils sont de la même famille et assez proches")
header(ws, r, ["Catégorie", "Famille", "Distance max (km) – actuelle", "Écart max (heures) – actuel", "Votre proposition (km / h)"])
for cat, (fam, km, h) in FAMILIES.items():
    r += 1
    put(ws, r, 1, model.CATEGORIES.get(cat, {}).get("fr", cat), bold=True)
    put(ws, r, 2, fam)
    put(ws, r, 3, km)
    put(ws, r, 4, h)
    put(ws, r, 5, None, "in")
r += 1
put(ws, r, 1, "Presse : mots communs minimaux entre deux titres", bold=True)
put(ws, r, 2, "Part de mots identiques (indice de Jaccard)")
put(ws, r, 3, 0.3, fmt="0%")
put(ws, r, 5, None, "in")
finish(ws, "A5")

# ================================================================== 13. Alerte précoce
ws = sheet("Alerte précoce", "Alerte précoce climat-conflit : barème",
           "Indice 0-100 par région administrative : climat (40 points max), sécurité (35), humanitaire (25). Niveau 1-4 selon le total. "
           "Modifiez les seuils et les points ; le simulateur en bas recalcule.",
           "Barème de l'indice d'alerte précoce et seuils des niveaux")
widths(ws, 46, 14, 14, 14, 14, 40)
E = {}
r = section(ws, 5, "Barème")
header(ws, r, ["Facteur", "Seuil actuel", "Votre seuil", "Points actuels", "Vos points", "Explication"])
bar = [
    ("rain1", "Pluie 3 mois / normale inférieure à", 0.5, 25, "Sécheresse sévère"), ("rain2", "Pluie / normale inférieure à", 0.7, 15, ""),
    ("rain3", "Pluie / normale inférieure à", 0.85, 7, "Déficit modéré"), ("rain4", "Pluie / normale supérieure à", 1.8, 15, "Pluies extrêmes"),
    ("rain5", "Pluie / normale supérieure à", 1.4, 8, "Excès de pluie"), ("temp1", "Écart de température (°C) supérieur à", 1.5, 10, "Chaleur anormale"),
    ("temp2", "Écart de température (°C) supérieur à", 0.8, 5, ""), ("soil1", "Humidité du sol sous le record (oui/non)", "oui", 5, ""),
    ("soil2", "Humidité du sol sous la normale (−0,1)", "oui", 3, ""), ("clim_cap", "Plafond climat", "", 40, ""),
    ("con1", "Incidents sur 3 mois au moins", 30, 20, ""), ("con2", "Incidents sur 3 mois au moins", 10, 14, ""),
    ("con3", "Incidents sur 3 mois au moins", 3, 8, ""), ("con4", "Incidents sur 3 mois au moins", 1, 3, ""),
    ("trend1", "Hausse vs normale (multiple) au moins", 1.5, 10, "Avec au moins 3 incidents"), ("trend2", "Hausse vs normale (multiple) au moins", 1.2, 5, ""),
    ("new", "Conflit nouveau (aucun antécédent, ≥ 3 incidents)", "", 8, ""), ("res", "Conflit lié aux ressources (eau, terres, bétail)", "", 5, ""),
    ("sec_cap", "Plafond sécurité", "", 35, ""), ("ipc1", "Part de population en IPC 3+ au moins", 0.4, 15, "Insécurité alimentaire"),
    ("ipc2", "Part en IPC 3+ au moins", 0.25, 10, ""), ("ipc3", "Part en IPC 3+ au moins", 0.15, 5, ""),
    ("idp1", "Déplacés / population au moins", 0.1, 10, ""), ("idp2", "Déplacés / population au moins", 0.03, 6, ""),
    ("idp3", "Déplacés présents (part plus faible)", "", 2, ""), ("hum_cap", "Plafond humanitaire", "", 25, ""),
    ("lvl4", "Niveau 4 (alerte) si total au moins", 60, "", ""), ("lvl3", "Niveau 3 si total au moins", 40, "", ""),
    ("lvl2", "Niveau 2 si total au moins", 20, "", "")]
for key, lab, th, pts, expl in bar:
    r += 1
    put(ws, r, 1, lab, bold=True)
    put(ws, r, 2, th, fmt="0%" if isinstance(th, float) and th < 1 and key.startswith(("ipc", "idp")) else None)
    put(ws, r, 3, th if th != "" else None, "in" if th != "" else "cur", "0%" if isinstance(th, float) and th < 1 and key.startswith(("ipc", "idp")) else None)
    put(ws, r, 4, pts)
    put(ws, r, 5, pts if pts != "" else None, "in" if pts != "" else "cur")
    put(ws, r, 6, expl, "note")
    E[key] = (f"$C${r}", f"$E${r}")
r = section(ws, r + 2, "Simulateur")
header(ws, r, ["Cas", "Pluie / normale", "Écart temp. (°C)", "Sol sous record (oui/non)", "Sol sous normale (oui/non)", "Incidents 3 mois"])
sim_rows = []
for lab, rain, t, s1, s2, n3 in [("Région sèche en conflit (type Darfour)", 0.45, 1.2, "non", "oui", 40),
                                  ("Région calme, pluies normales", 1.0, 0.3, "non", "non", 0), ("Votre cas", None, None, None, None, None)]:
    r += 1
    for i, v in enumerate([lab, rain, t, s1, s2, n3]):
        put(ws, r, i + 1, v, "in")
    sim_rows.append(r)
r += 1
header(ws, r, ["Cas (suite)", "Normale 3 mois (incidents)", "Conflit lié aux ressources (oui/non)", "Part IPC 3+", "Part de déplacés", "—"])
sim2 = []
for i, (base, res, ipc, idp) in enumerate([(12, "oui", 0.35, 0.12), (1, "non", 0, 0), (None, None, None, None)]):
    r += 1
    put(ws, r, 1, f"=A{sim_rows[i]}", "calc")
    for j, v in enumerate([base, res, ipc, idp]):
        put(ws, r, 2 + j, v, "in", "0%" if j >= 2 else None)
    sim2.append(r)
r += 1
header(ws, r, ["Résultat", "Climat (/40)", "Sécurité (/35)", "Humanitaire (/25)", "Total (/100)", "Niveau (1-4)"])
for i, a in enumerate(sim_rows):
    b = sim2[i]
    r += 1
    put(ws, r, 1, f"=A{a}", "calc")
    rain = (f"IF(B{a}=\"\",0,IF(B{a}<{E['rain1'][0]},{E['rain1'][1]},IF(B{a}<{E['rain2'][0]},{E['rain2'][1]},IF(B{a}<{E['rain3'][0]},{E['rain3'][1]},"
            f"IF(B{a}>{E['rain4'][0]},{E['rain4'][1]},IF(B{a}>{E['rain5'][0]},{E['rain5'][1]},0))))))")
    temp = f"IF(C{a}=\"\",0,IF(C{a}>{E['temp1'][0]},{E['temp1'][1]},IF(C{a}>{E['temp2'][0]},{E['temp2'][1]},0)))"
    soil = f"IF(D{a}=\"oui\",{E['soil1'][1]},IF(E{a}=\"oui\",{E['soil2'][1]},0))"
    put(ws, r, 2, f"=MIN({E['clim_cap'][1]},{rain}+{temp}+{soil})", "calc", "0")
    n = f"F{a}"
    cnt = (f"IF({n}>={E['con1'][0]},{E['con1'][1]},IF({n}>={E['con2'][0]},{E['con2'][1]},IF({n}>={E['con3'][0]},{E['con3'][1]},"
           f"IF({n}>={E['con4'][0]},{E['con4'][1]},0))))")
    trend = (f"IF({n}<3,0,IF(B{b}=0,{E['new'][1]},IF(B{b}=\"\",0,IF({n}/B{b}>={E['trend1'][0]},{E['trend1'][1]},"
             f"IF({n}/B{b}>={E['trend2'][0]},{E['trend2'][1]},0)))))")
    res = f"IF(C{b}=\"oui\",{E['res'][1]},0)"
    put(ws, r, 3, f"=MIN({E['sec_cap'][1]},{cnt}+{trend}+{res})", "calc", "0")
    ipc = (f"IF(D{b}=\"\",0,IF(D{b}>={E['ipc1'][0]},{E['ipc1'][1]},IF(D{b}>={E['ipc2'][0]},{E['ipc2'][1]},"
           f"IF(D{b}>={E['ipc3'][0]},{E['ipc3'][1]},0))))")
    idp = (f"IF(OR(E{b}=\"\",E{b}=0),0,IF(E{b}>={E['idp1'][0]},{E['idp1'][1]},IF(E{b}>={E['idp2'][0]},{E['idp2'][1]},{E['idp3'][1]})))")
    put(ws, r, 4, f"=MIN({E['hum_cap'][1]},{ipc}+{idp})", "calc", "0")
    put(ws, r, 5, f"=B{r}+C{r}+D{r}", "calc", "0", bold=True)
    put(ws, r, 6, f"=IF(E{r}>={E['lvl4'][0]},4,IF(E{r}>={E['lvl3'][0]},3,IF(E{r}>={E['lvl2'][0]},2,1)))", "calc", "0", bold=True)
finish(ws, "A5")

# ================================================================== 14. Go / no-go
ws = sheet("Go-no-go", "Go / no-go guidé (Travel buddy et brief de mission)",
           "Décision de déplacement : menace de la destination (M, 1-5) × vulnérabilité du voyageur (V, 1-5) = risque résiduel (1-25). "
           "Modifiez les ajustements de menace, les points de vulnérabilité et les seuils de décision ; testez avec le simulateur.",
           "Calcul de la menace, points de vulnérabilité, seuils de décision")
widths(ws, 40, 30, 12, 12, 12, 12, 44)
G = {}
r = section(ws, 5, "Ajustements de la menace (point de départ : niveau de risque du pays)")
header(ws, r, ["Règle", "Effet actuel", "Valeur", "Votre valeur", "", "", "Commentaire"])
for key, lab, eff, val in [("meae3", "MEAE niveau 3 (déconseillé sauf raison impérative)", "menace au moins", 4),
                           ("meae4", "MEAE niveau 4 (formellement déconseillé)", "menace au moins", 5),
                           ("us4", "US State Dept niveau 4", "menace au moins", 4),
                           ("pulse_low", "Pulse inférieur à", "+1", 35), ("pulse_drop", "Pulse en baisse sur 7 j d'au moins", "+1", 10),
                           ("near", "Incidents graves à moins de 100 km en 7 j, au moins", "+1", 3),
                           ("elec", "Élection pendant le séjour, si risque pays au moins", "+1", 3)]:
    r += 1
    put(ws, r, 1, lab, bold=True)
    put(ws, r, 2, eff)
    put(ws, r, 3, val)
    put(ws, r, 4, val, "in")
    put(ws, r, 7, None, "in")
    G[key] = f"$D${r}"
put(ws, r + 1, 1, "Zones partielles (MEAE « certaines zones ») : le niveau retenu est diminué de 1.", "note")
r = section(ws, r + 3, "Points de vulnérabilité (V = 1 + somme des points, arrondi, entre 1 et 5)")
header(ws, r, ["Question", "Réponse", "Si menace ≥", "Points", "Si menace ≥ (renforcé)", "Points renforcés", "Condition proposée au voyageur"])
VQ = [
    ("Profil", "Dirigeant exposé", 0, 1, 99, 1, ""), ("Profil", "Voyageuse", 3, 0.5, 99, 0.5, ""),
    ("Expérience", "Aucune", 0, 0.5, 3, 1, "Briefing approfondi, accueil par un correspondant local"),
    ("Appui sur place", "Aucun", 3, 1, 99, 1, "Prévoir un appui sur place"),
    ("Transport", "Conduite personnelle", 3, 1.5, 99, 1.5, "Chauffeur de confiance"),
    ("Transport", "Transports publics", 2, 0.5, 3, 1.5, "Chauffeur réservé"),
    ("Transport", "Taxi / VTC", 3, 0.5, 99, 0.5, "Chauffeur réservé plutôt que taxi"),
    ("Nuit / hors des villes", "Hors des villes", 0, 0.5, 3, 1.5, "Supprimer les trajets de nuit hors des villes"),
    ("Nuit / hors des villes", "En ville", 3, 0.5, 99, 0.5, "Limiter les déplacements de nuit"),
    ("Hébergement", "Location / chez l'habitant", 3, 1, 99, 1, "Hôtel évalué"),
    ("Hébergement", "Hôtel non évalué", 3, 0.5, 99, 0.5, "Faire évaluer l'hôtel"),
    ("Assurance-assistance", "Non", 0, 1.5, 99, 1.5, "Souscrire / vérifier l'assurance"),
    ("Communications", "Aucun plan", 0, 1, 99, 1, "Points de contact planifiés"),
    ("Communications", "Téléphone seulement", 4, 0.5, 99, 0.5, "Points de contact + téléphone satellite"),
    ("Préparation sûreté", "Aucune", 3, 0.5, 4, 1.5, "Briefing (menace 3) / formation HEAT (menace ≥ 4)"),
    ("Préparation sûreté", "Briefing", 5, 1, 99, 1, "Formation HEAT")]
vq_first = r + 1
for q in VQ:
    r += 1
    put(ws, r, 1, q[0], bold=True)
    put(ws, r, 2, q[1])
    for j, v in enumerate(q[2:6]):
        put(ws, r, 3 + j, v, "in")
    put(ws, r, 7, q[6], "in")
vq_last = r
put(ws, r + 1, 1, "99 = pas de palier renforcé. Les réponses absentes du tableau valent 0 point.", "note")
r = section(ws, r + 3, "Seuils de décision (risque résiduel R = M × V)")
header(ws, r, ["Décision", "R maximum – actuel", "Votre R maximum", "", "", "", "Commentaire"])
for key, lab, val in [("go", "GO – mesures standard", 6), ("cond", "GO sous conditions", 12), ("esc", "Escalade (décision de la direction sûreté)", 16)]:
    r += 1
    put(ws, r, 1, lab, bold=True)
    put(ws, r, 2, val)
    put(ws, r, 3, val, "in")
    put(ws, r, 7, None, "in")
    G[key] = f"$C${r}"
r += 1
put(ws, r, 1, "NO-GO", bold=True)
put(ws, r, 2, "au-delà")
put(ws, r, 7, "Règles spéciales : menace 5 → au moins Escalade ; menace 5 et V ≥ 3 → NO-GO ; mission reportable + Escalade/NO-GO → NO-GO ; "
    "mission vitale + NO-GO avec menace < 5 → Escalade.", "note")
r = section(ws, r + 2, "Simulateur")
OPTS = {"Criticité": ["Vitale", "Importante", "Reportable"],
        "Profil": ["Voyageur d'affaires", "Équipe technique", "Dirigeant exposé", "Voyageuse", "Séjour long"],
        "Expérience": ["Aucune", "Quelques séjours", "Habitué"], "Appui sur place": ["Aucun", "Partenaire", "Prestataire sûreté"],
        "Transport": ["Chauffeur réservé", "Taxi / VTC", "Conduite personnelle", "Transports publics"],
        "Nuit / hors des villes": ["Non", "En ville", "Hors des villes"],
        "Hébergement": ["Hôtel évalué", "Hôtel non évalué", "Location / chez l'habitant"], "Assurance-assistance": ["Oui", "Non"],
        "Communications": ["Complet", "Téléphone seulement", "Aucun plan"], "Préparation sûreté": ["Formation HEAT", "Briefing", "Aucune"]}
header(ws, r, ["Élément", "Votre saisie", "Points", "", "", "", "Explication"])
r += 1
put(ws, r, 1, "Niveau de risque du pays (1-5)", bold=True)
put(ws, r, 2, 3, "in")
lvl_cell = f"B{r}"
r += 1
put(ws, r, 1, "Avis MEAE (1-4, vide si aucun)", bold=True)
put(ws, r, 2, 3, "in")
meae_cell = f"B{r}"
r += 1
put(ws, r, 1, "MEAE : seulement certaines zones ? (oui/non)", bold=True)
put(ws, r, 2, "non", "in")
parts_cell = f"B{r}"
r += 1
put(ws, r, 1, "Avis US (1-4)", bold=True)
put(ws, r, 2, 3, "in")
us_cell = f"B{r}"
r += 1
put(ws, r, 1, "Pulse (0-100)", bold=True)
put(ws, r, 2, 55, "in")
pv_cell = f"B{r}"
r += 1
put(ws, r, 1, "Variation du Pulse sur 7 j", bold=True)
put(ws, r, 2, -4, "in")
pd_cell = f"B{r}"
r += 1
put(ws, r, 1, "Incidents graves à moins de 100 km (7 j)", bold=True)
put(ws, r, 2, 1, "in")
near_cell = f"B{r}"
r += 1
put(ws, r, 1, "Élection pendant le séjour ? (oui/non)", bold=True)
put(ws, r, 2, "non", "in")
el_cell = f"B{r}"
r += 1
put(ws, r, 1, "Menace M (1-5)", bold=True)
meae_eff = f'IF({meae_cell}="",0,IF({parts_cell}="oui",MAX(1,{meae_cell}-1),{meae_cell}))'
meae_map = f'IF({meae_eff}>=4,{G["meae4"]},IF({meae_eff}>=3,{G["meae3"]},{meae_eff}))'
base = f'MAX({lvl_cell},{meae_map},IF({us_cell}>=4,{G["us4"]},0))'
adds = (f'+IF({pv_cell}<{G["pulse_low"]},1,IF({pd_cell}<=-{G["pulse_drop"]},1,0))+IF({near_cell}>={G["near"]},1,0)'
        f'+IF(AND({el_cell}="oui",{lvl_cell}>={G["elec"]}),1,0)')
put(ws, r, 2, f"=MAX(1,MIN(5,{base}{adds}))", "calc", "0", bold=True)
T = f"B{r}"
answers = {}
for q, opts in OPTS.items():
    r += 1
    put(ws, r, 1, q, bold=True)
    # Exemple volontairement « moyen » : quelques points de vulnérabilité pour montrer l'effet sur la décision.
    default = {"Criticité": "Importante", "Profil": "Voyageuse", "Expérience": "Quelques séjours", "Appui sur place": "Partenaire",
               "Transport": "Taxi / VTC", "Nuit / hors des villes": "Non", "Hébergement": "Hôtel non évalué",
               "Assurance-assistance": "Oui", "Communications": "Téléphone seulement", "Préparation sûreté": "Briefing"}[q]
    put(ws, r, 2, default, "in")
    dv_list(ws, f"B{r}", opts)
    answers[q] = f"B{r}"
    if q != "Criticité":
        A_, B_, C_, D_, E_, F_ = [f"${c}${vq_first}:${c}${vq_last}" for c in "ABCDEF"]
        put(ws, r, 3, f'=SUMPRODUCT(({A_}=A{r})*({B_}=B{r})*({T}>={C_})*({D_}+({T}>={E_})*({F_}-{D_})))', "calc", "0.0")
pts_first, pts_last = r - len(OPTS) + 2, r
r += 1
put(ws, r, 1, "Vulnérabilité V (1-5)", bold=True)
put(ws, r, 2, f"=MAX(1,MIN(5,ROUND(1+SUM(C{pts_first}:C{pts_last}),0)))", "calc", "0", bold=True)
V = f"B{r}"
r += 1
put(ws, r, 1, "Risque résiduel R = M × V", bold=True)
put(ws, r, 2, f"={T}*{V}", "calc", "0", bold=True)
R = f"B{r}"
r += 1
put(ws, r, 1, "Décision brute", bold=True)
put(ws, r, 2, f'=IF({R}<={G["go"]},"GO",IF({R}<={G["cond"]},"GO sous conditions",IF({R}<={G["esc"]},"Escalade","NO-GO")))', "calc")
raw = f"B{r}"
r += 1
put(ws, r, 1, "Après règles « menace 5 »", bold=True)
put(ws, r, 2, f'=IF(AND({T}>=5,{V}>=3),"NO-GO",IF(AND({T}>=5,{raw}<>"NO-GO"),"Escalade",{raw}))', "calc")
t5 = f"B{r}"
r += 1
put(ws, r, 1, "DÉCISION FINALE (après criticité)", bold=True)
crit = answers["Criticité"]
put(ws, r, 2, f'=IF(AND({crit}="Reportable",OR({t5}="Escalade",{t5}="NO-GO")),"NO-GO",IF(AND({crit}="Vitale",{t5}="NO-GO",{T}<5),"Escalade",{t5}))', "calc", bold=True)
fin = r
ws.conditional_formatting.add(f"B{fin}", CellIsRule(operator="equal", formula=['"GO"'], fill=PatternFill("solid", fgColor="D4EDDA")))
ws.conditional_formatting.add(f"B{fin}", CellIsRule(operator="equal", formula=['"GO sous conditions"'], fill=PatternFill("solid", fgColor="FFF3CD")))
ws.conditional_formatting.add(f"B{fin}", CellIsRule(operator="equal", formula=['"Escalade"'], fill=PatternFill("solid", fgColor="FFE5CC")))
ws.conditional_formatting.add(f"B{fin}", CellIsRule(operator="equal", formula=['"NO-GO"'], fill=PatternFill("solid", fgColor="F8D7DA")))
finish(ws, "A5")

# ================================================================== 15. Recommandations
ws = sheet("Recommandations", "Recommandations du rapport pays",
           "Textes générés automatiquement dans le rapport pays selon le niveau du pays et les menaces élevées. C'est votre cœur de "
           "métier : réécrivez, complétez ou supprimez dans la colonne jaune.",
           "Postures par niveau et mesures par menace")
widths(ws, 30, 80, 80, 30)
r = 5
header(ws, r, ["Déclencheur", "Texte actuel", "Votre texte", "Commentaire"])
REC = [
    ("Niveau 1", "Politique voyage standard ; enregistrement des déplacements (Ariane pour les ressortissants français)."),
    ("Niveau 1", "Sensibilisation sûreté de base avant départ."),
    ("Niveau 2", "Politique voyage standard avec enregistrement systématique (Ariane) et suivi des voyageurs."),
    ("Niveau 2", "Briefing sûreté avant départ ; consignes de discrétion."),
    ("Niveau 3", "Validation hiérarchique des déplacements et briefing sûreté obligatoire."),
    ("Niveau 3", "Journey management : itinéraires validés, horaires, points de contact."),
    ("Niveau 3", "Hébergement dans des hôtels évalués ; transport avec chauffeur de confiance."),
    ("Niveau 4", "Déplacements limités aux besoins essentiels, validés par la direction sûreté."),
    ("Niveau 4", "Prestataire sûreté local, suivi GPS et points de contact réguliers."),
    ("Niveau 4", "Plan d'évacuation et de mise à l'abri testé ; assurance rapatriement et K&R."),
    ("Niveau 5", "Déplacements proscrits sauf mission vitale, sur décision de la direction."),
    ("Niveau 5", "Escorte armée ou véhicules protégés selon les zones ; mouvements de jour uniquement."),
    ("Niveau 5", "Plan d'évacuation actif, stocks de mise à l'abri (hibernation), communications satellitaires."),
    ("Niveau 5", "Revue quotidienne de la situation et critères de déclenchement de l'évacuation."),
    ("Menace Terrorisme ≥ 3", "Éviter les lieux de rassemblement, lieux de culte et sites symboliques lors des périodes sensibles."),
    ("Menace Terrorisme ≥ 3", "Varier itinéraires et horaires ; repérer les issues dans les lieux fréquentés."),
    ("Menace Conflit ≥ 3", "Cartographier les zones interdites et zones de front ; interdire tout déplacement non validé hors des zones autorisées."),
    ("Menace Conflit ≥ 3", "Abri renforcé sur site, protocole alerte aérienne / tirs indirects."),
    ("Menace Attaques/enlèvements ≥ 3", "Profil discret, pas de signes extérieurs de richesse ; transport sécurisé porte-à-porte."),
    ("Menace Attaques/enlèvements ≥ 3", "Sensibilisation anti-enlèvement ; assurance K&R et cellule de crise identifiée."),
    ("Menace Troubles ≥ 3", "Éviter les manifestations ; suivre les appels à la grève et à la mobilisation."),
    ("Menace Troubles ≥ 3", "Prévoir télétravail, stocks et itinéraires alternatifs pour les sites."),
    ("Menace Politique ≥ 3", "Suivre le calendrier politique (élections, votes, décisions judiciaires) et anticiper les couvre-feux."),
    ("Menace Politique ≥ 3", "Maintenir des liens avec l'ambassade et les réseaux d'entreprises locaux."),
    ("Menace Criminalité ≥ 3", "Sécurité physique des sites (contrôle d'accès, vidéosurveillance, gardiennage) adaptée au niveau de criminalité."),
    ("Menace Criminalité ≥ 3", "Limiter les déplacements de nuit ; distributeurs dans des lieux sûrs."),
    ("Menace Cyber ≥ 3", "Appareils dédiés pour les pays à risque, VPN, chiffrement ; vigilance face à l'hameçonnage."),
    ("Menace Catastrophes ≥ 3", "Plan de continuité d'activité tenant compte des saisons à risque (cyclones, moussons, feux)."),
    ("Menace Catastrophes ≥ 3", "Consignes séisme / inondation connues du personnel ; kits d'urgence sur site."),
    ("Menace Santé ≥ 3", "Vérifier vaccinations et prophylaxies avant départ ; assurance médicale et évacuation sanitaire."),
    ("Menace Santé ≥ 3", "Identifier les établissements de santé de référence près des sites."),
    ("Détention arbitraire citée", "Évaluer le profil des voyageurs (nationalité, binationaux, fonctions sensibles) avant validation."),
    ("Détention arbitraire citée", "Appareils « propres », aucun contenu sensible ; pas de photo de sites officiels ; contacts consulaires connus."),
    ("Tous pays (sites)", "Audit de sûreté des sites (accès, périmètre, protection des personnes) et plan de sûreté formalisé."),
    ("Tous pays (sites)", "Arbre d'alerte et exercices de gestion de crise au moins une fois par an."),
    ("Tous pays (sites)", "Veille quotidienne sur le pays et alertes automatiques dans un rayon défini autour de chaque site.")]
for trig, txt in REC:
    r += 1
    put(ws, r, 1, trig, bold=True)
    put(ws, r, 2, txt)
    put(ws, r, 3, None, "in")
    put(ws, r, 4, None, "in")
for _ in range(10):
    r += 1
    for col in range(1, 5):
        put(ws, r, col, None, "in")
finish(ws, "B6", f"A5:D{r}")

# ================================================================== 16. Villes
ws = sheet("Villes", "Notes d'analyste par ville (premier jet à valider)",
           f"{len(NOTES['cities'])} notes rédigées par Claude en septembre 2026 à partir de sources ouvertes. Elles s'affichent dans le "
           "rapport pays. Validez, corrigez le niveau (échelle Angor 1-5) et les textes, ou proposez d'autres villes (lignes en bas).",
           "Validation des notes par ville (niveau, résumé, zones, conseils, usages)")
widths(ws, 7, 18, 16, 8, 70, 50, 50, 40, 11, 10, 40)
r = 5
header(ws, r, ["Pays", "Nom du pays", "Ville", "Niveau", "Résumé", "Zones sensibles", "Conseils", "Us et coutumes", "Validé ?", "Votre niveau", "Vos corrections"])
for n in sorted(NOTES["cities"], key=lambda x: (CNAME.get(x["iso"], x["iso"]), x["city"])):
    r += 1
    put(ws, r, 1, n["iso"])
    put(ws, r, 2, CNAME.get(n["iso"], n["iso"]))
    put(ws, r, 3, n["city"], bold=True)
    put(ws, r, 4, n["level"])
    put(ws, r, 5, n["summary"])
    put(ws, r, 6, "\n".join("• " + z for z in n.get("zones", [])))
    put(ws, r, 7, "\n".join("• " + z for z in n.get("advice", [])))
    put(ws, r, 8, "\n".join("• " + z for z in n.get("customs", [])))
    put(ws, r, 9, None, "in")
    put(ws, r, 10, None, "in")
    put(ws, r, 11, None, "in")
last = r
for _ in range(15):
    r += 1
    for col in range(1, 12):
        put(ws, r, col, None, "in")
dv_list(ws, f"I6:I{r}", ["oui", "à corriger", "à supprimer"])
dv_list(ws, f"J6:J{r}", ["1", "2", "3", "4", "5"])
finish(ws, "D6", f"A5:K{last}")

# ================================================================== 17. Santé
ws = sheet("Santé", "Risques sanitaires du voyageur par pays",
           "Une ligne par pays, une colonne par risque. « x » = le risque s'applique (d'après OMS, CDC, HCSP, TravelHealthPro). "
           "Ajoutez ou retirez des « x » directement : c'est ce tableau qui alimente vaccins et maladies dans le rapport pays. "
           "« Rage faible » et « Hépatite A faible » : pays où ces risques sont considérés comme faibles.",
           "Pays concernés par chaque risque sanitaire")
LISTS = [("yellow_fever", "Fièvre jaune"), ("yellow_fever_cert", "Certificat FJ exigé"), ("malaria_high", "Paludisme étendu"),
         ("malaria_limited", "Paludisme limité"), ("dengue", "Dengue"), ("meningitis_belt", "Méningite (ceinture)"),
         ("japanese_encephalitis", "Encéphalite japonaise"), ("tb_high", "Tuberculose (forte charge)"), ("polio", "Polio sauvage"),
         ("cholera_recurrent", "Choléra récurrent"), ("altitude", "Altitude"), ("rabies_low", "Rage faible"), ("hep_a_low", "Hépatite A faible")]
widths(ws, 7, 26, *([11] * len(LISTS)), 36)
r = 5
header(ws, r, ["Pays", "Nom"] + [lab for _, lab in LISTS] + ["Commentaire"])
sets = {k: set(HEALTH["lists"].get(k, "").split()) for k, _ in LISTS}
for p in COUNTRIES:
    r += 1
    put(ws, r, 1, p["iso2"])
    put(ws, r, 2, p["name_fr"], bold=True)
    for j, (k, _) in enumerate(LISTS):
        c = put(ws, r, 3 + j, "x" if p["iso2"] in sets[k] else None, "in")
        c.alignment = CENTER
    put(ws, r, 3 + len(LISTS), None, "in")
r += 1
put(ws, r, 2, "Nombre de pays", bold=True)
for j in range(len(LISTS)):
    col = get_column_letter(3 + j)
    put(ws, r, 3 + j, f'=COUNTIF({col}6:{col}{r - 1},"x")', "calc", "0")
dv_list(ws, f"C6:{get_column_letter(2 + len(LISTS))}{r - 1}", ["x"])
finish(ws, "C6", f"A5:{get_column_letter(3 + len(LISTS))}{r - 1}")

# ================================================================== 18. Journal
ws = sheet("Journal", "Journal des décisions",
           "Notez ici chaque décision importante (une ligne par changement). Le développeur coche « Intégré » quand la décision est "
           "reportée dans l'outil, avec la version concernée.",
           "Traçabilité des décisions et de leur intégration")
widths(ws, 12, 22, 34, 24, 24, 44, 14, 12, 12)
r = 5
header(ws, r, ["Date", "Onglet", "Élément", "Ancienne valeur", "Nouvelle valeur", "Motif", "Auteur", "Intégré ?", "Version"])
r += 1
for col, v in enumerate([TODAY, "Note pays", "Poids des avis officiels", "0,6", "0,5",
                         "Exemple : les avis réagissent lentement ; donner plus de poids aux incidents récents", "Stef", "non", ""], 1):
    put(ws, r, col, v, "in")
for _ in range(60):
    r += 1
    for col in range(1, 10):
        put(ws, r, col, None, "in")
dv_list(ws, f"H6:H{r}", ["oui", "non", "rejeté"])
finish(ws, "A6", f"A5:I{r}")

# ================================================================== Lisez-moi
ws = readme
ws.sheet_view.showGridLines = False
widths(ws, 4, 28, 70, 18)
ws["B2"] = "Angor Intelligence – Revue de l'analyse des risques"
ws["B2"].font = Font(name=F, size=20, bold=True, color=NAVY)
ws["B3"] = f"Classeur généré le {TODAY} à partir de l'outil (version en cours). Tout ce qui est affiché correspond aux règles réellement appliquées aujourd'hui."
ws["B3"].font = SUB
r = 5
ws.cell(row=r, column=2, value="À quoi sert ce classeur").font = H2
for line in ["Relire et ajuster la façon dont l'outil classe, note et recommande : échelles, catégories, gravité, note pays, Pulse, "
             "cotation des sources, tri de la presse, alerte précoce, go / no-go, recommandations, notes par ville, santé.",
             "Les simulateurs (onglets Note pays, Matrice menaces, Pulse, Alerte précoce, Go-no-go) recalculent instantanément : "
             "changez un réglage jaune et voyez l'effet sur des cas concrets avant de décider.",
             "Rien ne change dans l'outil tant que le classeur n'est pas relu : renvoyez-le à Claude (ou à un développeur), qui reporte "
             "vos décisions dans la configuration et le code, puis coche « Intégré » dans l'onglet Journal."]:
    r += 1
    c = ws.cell(row=r, column=2, value="•  " + line)
    c.font, c.alignment = FONT, WRAP
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    ws.row_dimensions[r].height = 30
r += 2
ws.cell(row=r, column=2, value="Légende").font = H2
for fill, font, lab in [(INPUT_FILL, INPUT_FONT, "Jaune, texte bleu : cellule à modifier (proposition, valeur testée, votre avis)"),
                        (CALC_FILL, FONT, "Gris : calcul automatique (ne pas modifier)"),
                        (CUR_FILL, FONT, "Blanc : valeur actuelle de l'outil (référence, pour comparer)"),
                        (NOTE_FILL, Font(name=F, size=9, italic=True, color="555555"), "Beige : explication")]:
    r += 1
    c = ws.cell(row=r, column=2, value="Exemple")
    c.fill, c.font, c.border = fill, font, BOX
    ws.cell(row=r, column=3, value=lab).font = FONT
r += 2
ws.cell(row=r, column=2, value="Comment procéder").font = H2
for i, line in enumerate(["Commencez par les onglets de méthode (Échelles, Catégories, Note pays) : ils conditionnent tout le reste.",
                          "Dans chaque onglet, ne modifiez que les cellules jaunes. Laissez vide ce qui vous convient.",
                          "Utilisez les simulateurs pour vérifier qu'un réglage donne des niveaux cohérents avec votre expérience (colonne « Votre niveau attendu »).",
                          "Notez chaque décision importante dans l'onglet Journal (une ligne = un changement, avec le motif).",
                          "Suivez votre avancement dans la colonne « Statut » ci-dessous, puis renvoyez le classeur."], 1):
    r += 1
    c = ws.cell(row=r, column=2, value=f"{i}.  {line}")
    c.font, c.alignment = FONT, WRAP
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
    ws.row_dimensions[r].height = 28
r += 2
ws.cell(row=r, column=2, value="Onglets").font = H2
r += 1
header(ws, r, ["Onglet", "Ce que vous y décidez", "Statut"], col=2)
st_first = r + 1
for name, decide in TABS:
    r += 1
    c = ws.cell(row=r, column=2, value=name)
    c.hyperlink = Hyperlink(ref=c.coordinate, location=f"'{name}'!A1")   # lien interne (pas de cible externe)
    c.font, c.border = Font(name=F, size=10, color="1F5FD1", underline="single"), BOX
    put(ws, r, 3, decide)
    put(ws, r, 4, "à revoir", "in")
dv_list(ws, f"D{st_first}:D{r}", ["à revoir", "en cours", "terminé"])
ws.conditional_formatting.add(f"D{st_first}:D{r}", CellIsRule(operator="equal", formula=['"terminé"'], fill=PatternFill("solid", fgColor="D4EDDA")))
r += 1
put(ws, r, 3, "Onglets terminés", bold=True)
put(ws, r, 4, f'=COUNTIF(D{st_first}:D{r - 1},"terminé")&" / "&COUNTA(D{st_first}:D{r - 1})', "calc")
r += 2
ws.cell(row=r, column=2, value="Pour le développeur").font = H2
r += 1
c = ws.cell(row=r, column=2, value="Classeur produit par tools/build_risk_workbook.py (relancer pour repartir des valeurs actuelles). Correspondance : "
            "Note pays → config/risk.json · Catégories / Règles de gravité / Mots-clés → veille/model.py et veille/press.py · Pulse → veille/pulse.py · "
            "Cotation → veille/quality.py · Médias → tools/outlets_src.py · Sources → config/sources.json · Jeu de test → tests/gold_tri.csv · "
            "Fiabilité & alertes → docs/app.js, config/settings.json, veille/dedupe.py · Alerte précoce → veille/early_warning.py · "
            "Go-no-go → docs/gonogo.js · Recommandations → docs/report.js · Villes → config/city_notes.json · Santé → config/health.json.")
c.font, c.alignment = Font(name=F, size=9, color="555555"), WRAP
ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
ws.row_dimensions[r].height = 64
finish(ws)
ws.page_setup.orientation = "portrait"
wb.active = 0

if LANG == "en":
    from workbook_i18n import Translator, translate_workbook
    tr = Translator(json.loads((ROOT / "tools" / "risk_workbook_en.json").read_text(encoding="utf-8")),
                    {p["name_fr"]: p["name_en"] for p in COUNTRIES})
    translate_workbook(wb, tr)
    if tr.missing:
        where = ", ".join(f"{k} ({len(v)})" for k, v in tr.missing_in.items())
        print(f"{len(tr.missing)} textes sans traduction (laissés en français) : {where}")
    if "--missing" in sys.argv:
        (ROOT / "revue").mkdir(exist_ok=True)
        (ROOT / "revue" / "traductions_manquantes.json").write_text(
            json.dumps({k: "" for k in sorted(tr.missing)}, ensure_ascii=False, indent=1), encoding="utf-8")

for w in wb.worksheets:
    for row in w.iter_rows():
        for c in row:
            if c.font and c.font.name != F:
                c.font = Font(name=F, size=c.font.size, bold=c.font.bold, italic=c.font.italic, color=c.font.color)
wb.save(OUT)
print(f"Classeur écrit : {OUT} ({len(wb.sheetnames)} onglets)")
